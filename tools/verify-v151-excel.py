"""Read back real service and HTTP exports; never modify a workbook."""
from pathlib import Path
import json,datetime
from openpyxl import load_workbook
root=Path(__file__).resolve().parent.parent
books=[load_workbook(root/'output/audit-v151'/name,data_only=False) for name in ['receptions-export-test.xlsx','receptions-export-http.xlsx']]
results=[]
def check(name,predicate):
 assert all(predicate(w) for w in books),name
 results.append(dict(name=name,passed=True));print('PASS '+name)
check('Service et HTTP V1.5.1 : classeurs Excel ouverts sans réparation',lambda w:w.sheetnames==['Données','Contexte'])
check('Service et HTTP V1.5.1 : 213 lignes de données',lambda w:w['Données'].max_row==214)
check('Service et HTTP V1.5.1 : somme officielle 1809096 kg',lambda w:sum(w['Données'].cell(r,12).value or 0 for r in range(2,215))==1809096)
check('Service et HTTP V1.5.1 : dates Excel natives',lambda w:isinstance(w['Données']['D2'].value,datetime.datetime))
check('Service et HTTP V1.5.1 : en-tête figé et filtres',lambda w:w['Données'].freeze_panes=='A2' and w['Données'].auto_filter.ref=='A1:P214')
check('Service et HTTP V1.5.1 : aucune formule injectée',lambda w:not any(c.data_type=='f' for row in w['Données'] for c in row))
crates=load_workbook(root/'output/audit-v151/caisserie-producers-http.xlsx',data_only=False)
sheet=crates['Données'];rows=list(sheet.values)[1:]
def crate_check(name,ok):
 assert ok,name
 results.append(dict(name=name,passed=True));print('PASS '+name)
producer75=next(r for r in rows if r[2]==75)
crate_check('HTTP situation Excel : producteur 75 et six quantités exactes',producer75[5:11]==(5079,6660,1581,110,120,10))
total=next(r for r in rows if r[3]=='TOTAL')
crate_check('HTTP situation Excel : total caisses et palettes réconcilié',total[5:11]==(95208,136218,41010,2044,2338,294) and all(sum(r[c] for r in rows if r[3]!='TOTAL')==total[c] for c in range(5,11)))
crate_check('HTTP situation Excel : balances S moins E sur chaque ligne',all(r[7]==r[6]-r[5] and r[10]==r[9]-r[8] for r in rows))
crate_check('HTTP situation Excel : aucune formule et colonnes numériques',not any(c.data_type=='f' for row in sheet for c in row) and all(isinstance(r[c],(int,float)) for r in rows for c in range(5,11)))
(root/'docs/tests-v151-excel.json').write_text(json.dumps(dict(passed=len(results),failed=0,results=results),ensure_ascii=False,indent=2),encoding='utf-8')
