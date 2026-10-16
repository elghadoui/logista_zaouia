from pathlib import Path
from openpyxl import load_workbook
import json,datetime
root=Path(__file__).resolve().parent.parent
results=[]
def check(name,ok):
 assert ok,name
 results.append(dict(name=name,passed=True));print('PASS '+name)
for kind in ['service','http']:
 w=load_workbook(root/f'output/audit-v160/export-{kind}.xlsx',data_only=False);s=w['Données'];rows=list(s.values)[1:]
 check(kind+' : 4 lignes, 24 colonnes, totaux nets et colis',s.max_row==5 and s.max_column==24 and sum(r[15] for r in rows)==21158.4 and sum(r[14] for r in rows)==4773)
 check(kind+' : palettes des bons de quai, pas palettes créées',sum(r[13] for r in rows)==21)
 check(kind+' : poids net, commercial et pesé séparés',all(r[16] in (None,'') for r in rows) and sum(r[23] for r in rows)==21858 and sum(r[17] for r in rows)==25852.2)
 check(kind+' : dates natives, filtres et absence de formules',isinstance(rows[0][4],datetime.datetime) and s.auto_filter.ref=='A1:X5' and not any(c.data_type=='f' for row in s for c in row))
 context=' '.join(str(c.value or '') for row in w['Contexte'] for c in row)
 check(kind+' : définition et contexte traçables', '1.6.0' in context and 'TOTAL OFFICIEL FILTRÉ' in context and 'codmaj' in context)
w=load_workbook(root/'output/audit-v160/export-official.xlsx',data_only=False);s=w['Données'];rows=list(s.values)[1:]
check('Officiel : deux BDQ, un dossier, filtre codmaj 2',len(rows)==2 and len(set(r[2] for r in rows))==1 and all(r[5]==2 for r in rows))
check('Officiel : somme pdscom 7739 kg, 9 palettes, 3362 colis',sum(r[16] for r in rows)==7739 and sum(r[13] for r in rows)==9 and sum(r[14] for r in rows)==3362)
check('Officiel : client, statut et exportateur séparés',s.cell(1,7).value=='Client' and s.cell(1,9).value=='Exportateur' and all(r[18]=='Expédié et clôturé' for r in rows))
ctx=list(w['Contexte'].values);official=next(r for r in ctx if r[0]=='TOTAL OFFICIEL FILTRÉ kg')
check('Officiel : récapitulatif exactement égal aux lignes filtrées',official[1]==sum(r[16] for r in rows) and official[3]==9 and official[5]==3362 and official[7]==1)
(root/'docs/tests-v160-excel.json').write_text(json.dumps(dict(passed=len(results),failed=0,results=results),ensure_ascii=False,indent=2),encoding='utf-8')
