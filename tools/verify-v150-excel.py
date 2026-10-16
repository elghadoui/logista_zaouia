"""Read back real service and HTTP exports; never modify a workbook."""
from pathlib import Path
import json,datetime
from openpyxl import load_workbook
root=Path(__file__).resolve().parent.parent
books=[load_workbook(root/'output/audit-v150'/name,data_only=False) for name in ['receptions-export-test.xlsx','receptions-export-http.xlsx']]
results=[]
def check(name,predicate):
 assert all(predicate(w) for w in books),name
 results.append(dict(name=name,passed=True));print('PASS '+name)
check('Service et HTTP : classeurs Excel ouverts sans réparation',lambda w:w.sheetnames==['Données','Contexte'])
check('Service et HTTP : 213 lignes de données',lambda w:w['Données'].max_row==214)
check('Service et HTTP : somme officielle 1809096 kg',lambda w:sum(w['Données'].cell(r,12).value or 0 for r in range(2,215))==1809096)
check('Service et HTTP : dates Excel natives',lambda w:isinstance(w['Données']['D2'].value,datetime.datetime))
check('Service et HTTP : en-tête figé et filtres',lambda w:w['Données'].freeze_panes=='A2' and w['Données'].auto_filter.ref=='A1:P214')
check('Service et HTTP : aucune formule injectée',lambda w:not any(c.data_type=='f' for row in w['Données'] for c in row))
(root/'docs/tests-v150-excel.json').write_text(json.dumps(dict(passed=len(results),failed=0,results=results),ensure_ascii=False,indent=2),encoding='utf-8')
