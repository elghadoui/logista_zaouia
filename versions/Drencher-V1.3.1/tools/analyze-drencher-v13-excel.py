import json, hashlib, pathlib, datetime
import openpyxl
paths = [pathlib.Path(r'C:\Users\abdellah\Desktop\DeskTops\drensher\F04 MPS-01 Suivi du traitement au niveau du drencher V05 (2).xlsx'),pathlib.Path(r"C:\Users\abdellah\Desktop\l'état Clémentine Oranges 2026 2026.xlsx")]
result=[]
for path in paths:
 w=openpyxl.load_workbook(path,data_only=False)
 cached=openpyxl.load_workbook(path,data_only=True)
 sheets=[]
 for s in w:
  cells=[{'cell':c.coordinate,'value':str(c.value),'cached':str(cached[s.title][c.coordinate].value),'format':c.number_format} for row in s for c in row if c.value is not None]
  sheets.append({'name':s.title,'dimension':s.calculate_dimension(),'printArea':str(s.print_area),'printTitles':s.print_titles,'orientation':s.page_setup.orientation,'paper':s.page_setup.paperSize,'fitWidth':s.page_setup.fitToWidth,'fitHeight':s.page_setup.fitToHeight,'merged':[str(m) for m in s.merged_cells],'tables':list(s.tables),'images':len(s._images),'rowGroups':[{ 'row':i,'level':v.outlineLevel,'hidden':v.hidden} for i,v in s.row_dimensions.items() if v.outlineLevel],'cells':cells})
 result.append({'path':str(path),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'sheets':sheets})
out=pathlib.Path('docs/drencher-v13-excel-analysis.json');out.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
for item in result:
 print(item['path'])
 for s in item['sheets']:
  print({k:v for k,v in s.items() if k!='cells'})
  print(json.dumps(s['cells'][:85],ensure_ascii=False))
